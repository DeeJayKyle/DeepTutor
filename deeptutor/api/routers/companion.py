"""Tutor Companion API: AI teaching materials, research lookups, and backups.

The Tutor Companion web app (``web/public/companion``) keeps every class,
student and mark in the tutor's own browser storage. This router only adds the
parts that need a server:

* ``/generate`` — drafts lecture notes, exams, quizzes, assignments and slide
  outlines with the server-configured LLM, so no API key reaches the client.
* ``/extract`` — pulls plain text out of uploaded textbooks and curricula so
  the client can cache it and pass excerpts to ``/generate``.
* ``/research`` — looks up articles (OpenAlex, arXiv), openly licensed images
  (Wikimedia Commons) and YouTube videos. Every result comes from a real index
  and carries its source URL; nothing here is written by the LLM.
* ``/backups`` — optional, tutor-initiated backups stored in the signed-in
  user's own workspace, versioned and checked with SHA-256.
"""

from __future__ import annotations

import asyncio
from datetime import datetime, timezone
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import re
from typing import Any, Literal
import urllib.parse
import uuid
import xml.etree.ElementTree as ET

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
import httpx
from pydantic import BaseModel, Field

from deeptutor.api.routers.auth import require_auth
from deeptutor.multi_user.paths import get_current_path_service

logger = logging.getLogger(__name__)

router = APIRouter()

MAX_SOURCE_CHARS = 60_000
MAX_EXTRACT_BYTES = 40 * 1024 * 1024
MAX_EXTRACT_CHARS = 400_000
MAX_BACKUP_BYTES = 60 * 1024 * 1024
MAX_BACKUPS_KEPT = 20
HTTP_TIMEOUT = 15.0
USER_AGENT = "TutorCompanion/1.0 (+https://github.com/HKUDS/DeepTutor)"

Kind = Literal[
    "lecture_notes",
    "topic_summary",
    "exam",
    "quiz",
    "assignment",
    "presentation",
    "marking_guide",
    "case_study",
    "revise",
]

KIND_BRIEFS: dict[str, str] = {
    "lecture_notes": (
        "Write comprehensive lecture notes. Include: learning objectives, key definitions, "
        "detailed explanations organised by topic and subtopic, worked examples or case "
        "studies, a table where it aids comparison, practical exercises, revision "
        "questions, a summary of key learning points, and suggested further reading."
    ),
    "topic_summary": (
        "Write a concise topic summary: learning objectives, key points, essential "
        "definitions, and five revision questions."
    ),
    "exam": (
        "Write a complete examination paper. Start with a header (title, course, duration, "
        "total marks) and the instructions to candidates. Number every question and show "
        "its marks in brackets, e.g. '(5 marks)'. Group questions into sections by type. "
        "Make the marks add up exactly to the requested total."
    ),
    "quiz": (
        "Write a quiz. Number every question and show its marks. For multiple-choice "
        "questions list options A-D on separate lines."
    ),
    "assignment": (
        "Write an assignment brief: title, type, learning outcomes assessed, background, "
        "the task(s) with clear deliverables, submission requirements, and a marking "
        "rubric as a markdown table (criteria x performance levels with marks)."
    ),
    "presentation": (
        "Write a teaching slide deck. Each slide starts with a level-2 heading "
        "('## Slide title') followed by 3-6 short bullet points. Begin with a title slide "
        "and a learning-objectives slide; end with discussion questions, revision "
        "questions, and references slides. Add speaker notes under a line 'Notes:' when "
        "helpful."
    ),
    "marking_guide": (
        "Write a marking guide for the assessment provided in the instructions: for each "
        "question give the model answer, the marks allocation point by point, and what "
        "earns partial credit."
    ),
    "case_study": (
        "Write a realistic teaching case study with background, the scenario, data or "
        "observations, and discussion questions, followed by guidance notes for the tutor."
    ),
    "revise": (
        "Revise the existing draft given in the instructions exactly as the tutor asks. "
        "Return the complete revised document."
    ),
}

ANSWER_KINDS = {"exam", "quiz"}

SYSTEM_PROMPT = """You are Tutor Companion, an assistant that helps tutors prepare teaching \
and assessment materials. Write in clear, accurate, well-structured Markdown (headings, \
numbered lists, tables). Rules you must follow:
1. Treat the SOURCE MATERIALS supplied by the tutor as the primary authority. When a \
statement comes from a source, end it with a source tag like [Source: S1].
2. Anything you add from general knowledge must not carry a source tag. Put a short \
'> AI-generated explanation' blockquote marker before sections that are mainly your own \
elaboration.
3. Never invent citations, page numbers, URLs, statistics or authors. In any reference \
list, only list the supplied sources; you may suggest further reading as \
'Suggested topics to search' without fabricated bibliographic details.
4. Stay within the curriculum, level and topics requested.
5. The tutor will review your draft before use, so flag uncertain points with \
'(verify)'."""


class SourceExcerpt(BaseModel):
    title: str = Field(default="", max_length=300)
    text: str = Field(default="", max_length=MAX_SOURCE_CHARS)


class GenerateRequest(BaseModel):
    kind: Kind
    class_name: str = Field(default="", max_length=200)
    level: str = Field(default="", max_length=200)
    unit_title: str = Field(default="", max_length=200)
    unit_code: str = Field(default="", max_length=60)
    topics: list[str] = Field(default_factory=list, max_length=60)
    objectives: list[str] = Field(default_factory=list, max_length=60)
    curriculum: str = Field(default="", max_length=20_000)
    sources: list[SourceExcerpt] = Field(default_factory=list, max_length=12)
    instructions: str = Field(default="", max_length=60_000)
    detail: Literal["brief", "standard", "comprehensive"] = "standard"
    difficulty: Literal["introductory", "intermediate", "advanced", "mixed"] = "mixed"
    question_types: list[str] = Field(default_factory=list, max_length=12)
    num_questions: int | None = Field(default=None, ge=1, le=200)
    total_marks: int | None = Field(default=None, ge=1, le=1000)
    duration: str = Field(default="", max_length=60)
    slides: int | None = Field(default=None, ge=3, le=60)
    include_answers: bool = True


class ResearchRequest(BaseModel):
    query: str = Field(min_length=2, max_length=300)
    kind: Literal["articles", "images", "videos"] = "articles"
    limit: int = Field(default=8, ge=1, le=20)


class BackupCreate(BaseModel):
    label: str = Field(default="", max_length=120)
    device: str = Field(default="", max_length=120)
    sha256: str = Field(min_length=64, max_length=64)
    payload: str


# ---------------------------------------------------------------------------
# Status
# ---------------------------------------------------------------------------


def _ai_configured() -> tuple[bool, str]:
    try:
        from deeptutor.services.llm.config import get_llm_config

        config = get_llm_config()
    except Exception as exc:  # noqa: BLE001 — any resolver failure means "not configured"
        logger.debug("Companion: LLM config unavailable: %s", exc)
        return False, ""
    return bool(config.model), str(config.model or "")


@router.get("/status")
async def status(current=Depends(require_auth)) -> dict[str, Any]:
    ai_ready, model = _ai_configured()
    return {
        "user": getattr(current, "username", "") or "local",
        "role": getattr(current, "role", "") or "admin",
        "ai_available": ai_ready,
        "model": model,
        "youtube_api": bool(os.environ.get("YOUTUBE_API_KEY")),
        "server_time": datetime.now(timezone.utc).isoformat(),
    }


# ---------------------------------------------------------------------------
# Generation
# ---------------------------------------------------------------------------


def build_prompt(req: GenerateRequest) -> str:
    """Assemble the user prompt for a generation request (pure; unit-tested)."""
    lines: list[str] = [f"TASK: {KIND_BRIEFS[req.kind]}", ""]
    context = [
        ("Class", req.class_name),
        ("Academic level", req.level),
        ("Course unit", " ".join(p for p in (req.unit_code, req.unit_title) if p)),
        ("Level of detail", req.detail),
        ("Difficulty", req.difficulty),
    ]
    lines += [f"{label}: {value}" for label, value in context if value]
    if req.topics:
        lines.append("Topics to cover: " + "; ".join(t.strip() for t in req.topics if t.strip()))
    if req.objectives:
        lines.append("Learning outcomes:")
        lines += [f"- {o.strip()}" for o in req.objectives if o.strip()]
    if req.kind in ANSWER_KINDS or req.kind == "assignment":
        if req.question_types:
            lines.append("Question types: " + ", ".join(req.question_types))
        if req.num_questions:
            lines.append(f"Number of questions: {req.num_questions}")
        if req.total_marks:
            lines.append(f"Total marks: {req.total_marks}")
        if req.duration:
            lines.append(f"Duration: {req.duration}")
    if req.kind in ANSWER_KINDS:
        if req.include_answers:
            lines.append(
                "After the paper, add a section '# Marking Guide' with the model answer, "
                "marks allocation, explanation of the correct answer, and the learning "
                "outcome assessed for every question."
            )
        else:
            lines.append("Do not include answers.")
    if req.kind == "presentation" and req.slides:
        lines.append(f"Number of slides: about {req.slides}")
    if req.curriculum.strip():
        lines += ["", "CURRICULUM / COURSE OUTLINE:", req.curriculum.strip()]
    if req.sources:
        lines += ["", "SOURCE MATERIALS (cite as [Source: S#]):"]
        for index, source in enumerate(req.sources, start=1):
            title = source.title.strip() or f"Source {index}"
            lines += [f"--- S{index}: {title} ---", source.text.strip()]
    else:
        lines += [
            "",
            "No source materials were supplied: mark the whole draft as AI-generated and "
            "do not use [Source] tags.",
        ]
    if req.instructions.strip():
        lines += ["", "TUTOR INSTRUCTIONS:", req.instructions.strip()]
    return "\n".join(lines)


@router.post("/generate")
async def generate(req: GenerateRequest, _current=Depends(require_auth)) -> dict[str, Any]:
    ai_ready, model = _ai_configured()
    if not ai_ready:
        raise HTTPException(
            status_code=503,
            detail="No AI model is configured on the server. Ask the administrator to add one in Settings.",
        )
    from deeptutor.services.llm import complete

    prompt = build_prompt(req)
    try:
        content = await complete(prompt, system_prompt=SYSTEM_PROMPT)
    except Exception as exc:  # noqa: BLE001 — surface provider failures as a clean 502
        logger.warning("Companion generation failed: %s", exc)
        raise HTTPException(status_code=502, detail=f"The AI service failed: {exc}") from exc
    content = (content or "").strip()
    if not content:
        raise HTTPException(status_code=502, detail="The AI service returned an empty response.")
    return {
        "content": content,
        "model": model,
        "sources_used": [s.title or f"Source {i}" for i, s in enumerate(req.sources, start=1)],
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "ai_generated": True,
    }


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------


def _extract_pdf(data: bytes) -> str:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        return "\n\n".join(page.get_text() for page in doc)


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for number, slide in enumerate(deck.slides, start=1):
        parts.append(f"[Slide {number}]")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text_frame.text)
    return "\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in book.worksheets:
        parts.append(f"[Sheet {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
}
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".html", ".htm", ".json"}


@router.post("/extract")
async def extract(file: UploadFile = File(...), _current=Depends(require_auth)) -> dict[str, Any]:
    name = file.filename or "upload"
    suffix = Path(name).suffix.lower()
    data = await file.read(MAX_EXTRACT_BYTES + 1)
    if len(data) > MAX_EXTRACT_BYTES:
        raise HTTPException(status_code=413, detail="File is larger than 40 MB.")
    if not data:
        raise HTTPException(status_code=400, detail="The file is empty.")
    if suffix in TEXT_SUFFIXES:
        text = data.decode("utf-8", errors="replace")
    elif suffix in EXTRACTORS:
        try:
            text = await asyncio.to_thread(EXTRACTORS[suffix], data)
        except Exception as exc:  # noqa: BLE001 — corrupted files are a user error, not a crash
            raise HTTPException(
                status_code=422, detail=f"Could not read {name}: the file may be corrupted ({exc})."
            ) from exc
    else:
        raise HTTPException(
            status_code=415,
            detail=f"Text extraction is not supported for '{suffix or name}' files.",
        )
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    truncated = len(text) > MAX_EXTRACT_CHARS
    return {
        "filename": name,
        "chars": len(text),
        "truncated": truncated,
        "text": text[:MAX_EXTRACT_CHARS],
    }


# ---------------------------------------------------------------------------
# Research (real indexes only — no LLM-written references)
# ---------------------------------------------------------------------------


def _openalex_abstract(index: dict[str, list[int]] | None) -> str:
    if not index:
        return ""
    positions = sorted((pos, word) for word, spots in index.items() for pos in spots)
    return " ".join(word for _, word in positions)[:600]


def parse_openalex(payload: dict[str, Any]) -> list[dict[str, Any]]:
    results = []
    for work in payload.get("results", []):
        doi = work.get("doi") or ""
        location = work.get("primary_location") or {}
        source = (location.get("source") or {}).get("display_name") or ""
        open_access = work.get("open_access") or {}
        url = (
            doi
            or open_access.get("oa_url")
            or location.get("landing_page_url")
            or work.get("id", "")
        )
        results.append(
            {
                "title": work.get("display_name") or work.get("title") or "Untitled",
                "authors": [
                    (a.get("author") or {}).get("display_name", "")
                    for a in (work.get("authorships") or [])[:6]
                ],
                "date": work.get("publication_date") or str(work.get("publication_year") or ""),
                "source": source or "OpenAlex",
                "url": url,
                "doi": doi,
                "open_access_url": open_access.get("oa_url") or "",
                "cited_by": work.get("cited_by_count", 0),
                "summary": _openalex_abstract(work.get("abstract_inverted_index")),
                "provider": "OpenAlex",
            }
        )
    return results


def parse_arxiv(xml_text: str) -> list[dict[str, Any]]:
    ns = {"a": "http://www.w3.org/2005/Atom"}
    root = ET.fromstring(xml_text)
    results = []
    for entry in root.findall("a:entry", ns):

        def text(tag: str) -> str:
            node = entry.find(tag, ns)
            return " ".join((node.text or "").split()) if node is not None else ""

        results.append(
            {
                "title": text("a:title"),
                "authors": [
                    " ".join((n.text or "").split()) for n in entry.findall("a:author/a:name", ns)
                ][:6],
                "date": text("a:published")[:10],
                "source": "arXiv (preprint, not peer reviewed)",
                "url": text("a:id"),
                "doi": "",
                "open_access_url": text("a:id"),
                "cited_by": None,
                "summary": text("a:summary")[:600],
                "provider": "arXiv",
            }
        )
    return results


_TAG_RE = re.compile(r"<[^>]+>")


def parse_commons(payload: dict[str, Any]) -> list[dict[str, Any]]:
    pages = (payload.get("query") or {}).get("pages") or {}
    results = []
    for page in sorted(pages.values(), key=lambda p: p.get("index", 0)):
        info = (page.get("imageinfo") or [{}])[0]
        meta = info.get("extmetadata") or {}

        def field(name: str) -> str:
            return _TAG_RE.sub("", (meta.get(name) or {}).get("value", "")).strip()

        results.append(
            {
                "title": page.get("title", "").removeprefix("File:"),
                "thumbnail": info.get("thumburl") or info.get("url", ""),
                "url": info.get("descriptionurl") or info.get("url", ""),
                "file_url": info.get("url", ""),
                "author": field("Artist"),
                "license": field("LicenseShortName"),
                "license_url": (meta.get("LicenseUrl") or {}).get("value", ""),
                "attribution_required": field("AttributionRequired") == "true",
                "provider": "Wikimedia Commons",
            }
        )
    return results


def parse_youtube(payload: dict[str, Any]) -> list[dict[str, Any]]:
    results = []
    for item in payload.get("items", []):
        video_id = (item.get("id") or {}).get("videoId")
        if not video_id:
            continue
        snippet = item.get("snippet") or {}
        results.append(
            {
                "title": snippet.get("title", ""),
                "channel": snippet.get("channelTitle", ""),
                "date": (snippet.get("publishedAt") or "")[:10],
                "url": f"https://www.youtube.com/watch?v={video_id}",
                "thumbnail": ((snippet.get("thumbnails") or {}).get("medium") or {}).get("url", ""),
                "summary": snippet.get("description", ""),
                "provider": "YouTube",
            }
        )
    return results


async def _get(client: httpx.AsyncClient, url: str, params: dict[str, Any]) -> httpx.Response:
    response = await client.get(url, params=params)
    response.raise_for_status()
    return response


@router.post("/research")
async def research(req: ResearchRequest, _current=Depends(require_auth)) -> dict[str, Any]:
    query = req.query.strip()
    notes: list[str] = []
    results: list[dict[str, Any]] = []
    async with httpx.AsyncClient(
        timeout=HTTP_TIMEOUT, headers={"User-Agent": USER_AGENT}, follow_redirects=True
    ) as client:
        try:
            if req.kind == "articles":
                openalex, arxiv = await asyncio.gather(
                    _get(
                        client,
                        "https://api.openalex.org/works",
                        {"search": query, "per-page": req.limit},
                    ),
                    _get(
                        client,
                        "https://export.arxiv.org/api/query",
                        {"search_query": f"all:{query}", "max_results": max(2, req.limit // 3)},
                    ),
                    return_exceptions=True,
                )
                if isinstance(openalex, httpx.Response):
                    results += parse_openalex(openalex.json())
                else:
                    notes.append(f"OpenAlex unavailable: {openalex}")
                if isinstance(arxiv, httpx.Response):
                    results += parse_arxiv(arxiv.text)
                else:
                    notes.append(f"arXiv unavailable: {arxiv}")
            elif req.kind == "images":
                response = await _get(
                    client,
                    "https://commons.wikimedia.org/w/api.php",
                    {
                        "action": "query",
                        "format": "json",
                        "generator": "search",
                        "gsrsearch": f"filetype:bitmap|drawing {query}",
                        "gsrnamespace": 6,
                        "gsrlimit": req.limit,
                        "prop": "imageinfo",
                        "iiprop": "url|extmetadata",
                        "iiurlwidth": 320,
                    },
                )
                results = parse_commons(response.json())
            else:
                key = os.environ.get("YOUTUBE_API_KEY", "").strip()
                if key:
                    response = await _get(
                        client,
                        "https://www.googleapis.com/youtube/v3/search",
                        {
                            "part": "snippet",
                            "type": "video",
                            "q": query,
                            "maxResults": req.limit,
                            "safeSearch": "strict",
                            "key": key,
                        },
                    )
                    results = parse_youtube(response.json())
                else:
                    notes.append(
                        "YouTube search needs a YOUTUBE_API_KEY on the server. "
                        "Open the YouTube search link instead and save the videos you choose."
                    )
        except (httpx.HTTPError, ValueError, ET.ParseError) as exc:
            raise HTTPException(status_code=502, detail=f"Research service failed: {exc}") from exc
    return {
        "query": query,
        "kind": req.kind,
        "results": results,
        "notes": notes,
        "search_url": "https://www.youtube.com/results?search_query="
        + urllib.parse.quote_plus(query)
        if req.kind == "videos"
        else "",
        "verified_from": "Results are returned by public indexes and link to the original record.",
    }


# ---------------------------------------------------------------------------
# Backups (stored in the signed-in user's own workspace)
# ---------------------------------------------------------------------------

_BACKUP_ID = re.compile(r"^[0-9]{8}T[0-9]{6}Z-[0-9a-f]{8}$")


def _backup_dir() -> Path:
    path = get_current_path_service().get_user_root() / "companion" / "backups"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _backup_path(backup_id: str) -> Path:
    if not _BACKUP_ID.match(backup_id):
        raise HTTPException(status_code=400, detail="Invalid backup id.")
    return _backup_dir() / f"{backup_id}.json"


def _summary(record: dict[str, Any]) -> dict[str, Any]:
    return {
        k: record[k] for k in ("id", "label", "device", "created_at", "size", "sha256", "counts")
    }


@router.get("/backups")
async def list_backups(_current=Depends(require_auth)) -> dict[str, Any]:
    items = []
    for path in sorted(_backup_dir().glob("*.json"), reverse=True):
        try:
            items.append(_summary(json.loads(path.read_text("utf-8"))))
        except (OSError, ValueError, KeyError):
            logger.warning("Skipping unreadable companion backup %s", path.name)
    return {"backups": items}


@router.post("/backups", status_code=201)
async def create_backup(body: BackupCreate, _current=Depends(require_auth)) -> dict[str, Any]:
    raw = body.payload.encode("utf-8")
    if len(raw) > MAX_BACKUP_BYTES:
        raise HTTPException(
            status_code=413,
            detail="Backup is larger than 60 MB. Exclude library files and try again.",
        )
    digest = hashlib.sha256(raw).hexdigest()
    if digest != body.sha256.lower():
        raise HTTPException(
            status_code=400, detail="Integrity check failed: the backup was corrupted in transit."
        )
    try:
        data = json.loads(body.payload)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Backup payload is not valid JSON.") from exc
    stores = data.get("stores") if isinstance(data, dict) else None
    if not isinstance(stores, dict):
        raise HTTPException(status_code=400, detail="Backup payload is missing its data stores.")
    now = datetime.now(timezone.utc)
    backup_id = f"{now.strftime('%Y%m%dT%H%M%SZ')}-{uuid.uuid4().hex[:8]}"
    record = {
        "id": backup_id,
        "label": body.label.strip() or "Backup",
        "device": body.device.strip(),
        "created_at": now.isoformat(),
        "size": len(raw),
        "sha256": digest,
        "counts": {k: len(v) for k, v in stores.items() if isinstance(v, list)},
        "payload": body.payload,
    }
    target = _backup_path(backup_id)
    temp = target.with_suffix(".tmp")
    temp.write_text(json.dumps(record), "utf-8")
    temp.replace(target)
    for old in sorted(_backup_dir().glob("*.json"), reverse=True)[MAX_BACKUPS_KEPT:]:
        old.unlink(missing_ok=True)
    return _summary(record)


@router.get("/backups/{backup_id}")
async def get_backup(backup_id: str, _current=Depends(require_auth)) -> dict[str, Any]:
    path = _backup_path(backup_id)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Backup not found.")
    record = json.loads(path.read_text("utf-8"))
    if hashlib.sha256(record["payload"].encode("utf-8")).hexdigest() != record["sha256"]:
        raise HTTPException(status_code=409, detail="Stored backup failed its integrity check.")
    return record


@router.delete("/backups/{backup_id}")
async def delete_backup(backup_id: str, _current=Depends(require_auth)) -> dict[str, Any]:
    path = _backup_path(backup_id)
    if not path.exists():
        raise HTTPException(status_code=404, detail="Backup not found.")
    path.unlink()
    return {"deleted": backup_id}
